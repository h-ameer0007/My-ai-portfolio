"""DOCX export functionality for YouTube Semantic Optimizer."""

import os
from datetime import datetime
from typing import Dict, Optional
from pathlib import Path

from config.settings import EXPORTS_DIR, STEPS


def export_step_output(project_name: str, step_number: int,
                       output: str) -> str:
    """Export a single step output as a text file."""
    safe_name = "".join(c if c.isalnum() or c in " -_" else "" for c in project_name)
    filename = f"{safe_name}_Step{step_number}_{STEPS.get(step_number, 'Unknown')}.txt"
    filename = filename.replace(" ", "_")
    filepath = EXPORTS_DIR / filename

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f"YouTube Semantic Optimizer - Step {step_number}\n")
        f.write(f"Project: {project_name}\n")
        f.write(f"Step: {STEPS.get(step_number, 'Unknown')}\n")
        f.write(f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write("=" * 60 + "\n\n")
        f.write(output)

    return str(filepath)


def export_full_report_docx(project_name: str, project_data: Dict,
                            step_outputs: Dict[int, str]) -> str:
    """Export complete optimization report as DOCX."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return export_full_report_txt(project_name, project_data, step_outputs)

    doc = Document()

    # Title
    title = doc.add_heading("YouTube Semantic Optimizer Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Project info
    doc.add_heading("Project Information", level=1)
    info_table = doc.add_table(rows=4, cols=2)
    info_table.style = "Light List"
    info_data = [
        ("Project Name", project_name),
        ("Topic", project_data.get("topic", "")),
        ("Channel", project_data.get("channel", "")),
        ("Target Audience", project_data.get("target_audience", "")),
    ]
    for i, (label, value) in enumerate(info_data):
        info_table.cell(i, 0).text = label
        info_table.cell(i, 1).text = value

    doc.add_paragraph("")

    # Step outputs
    for step_num in sorted(step_outputs.keys()):
        output = step_outputs[step_num]
        if not output:
            continue

        step_name = STEPS.get(step_num, f"Step {step_num}")
        doc.add_heading(f"Step {step_num}: {step_name}", level=1)

        # Add output text, splitting into paragraphs
        for paragraph in output.split("\n"):
            if paragraph.strip():
                if paragraph.startswith("="):
                    doc.add_paragraph("_" * 50)
                else:
                    p = doc.add_paragraph(paragraph)
                    p.style.font.size = Pt(10)

    # Footer
    doc.add_paragraph("")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    safe_name = "".join(c if c.isalnum() or c in " -_" else "" for c in project_name)
    filename = f"{safe_name}_Full_Report.docx"
    filepath = EXPORTS_DIR / filename

    doc.save(str(filepath))
    return str(filepath)


def export_script_docx(project_name: str, script: str,
                       title: str = "") -> str:
    """Export script only as DOCX."""
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return export_script_txt(project_name, script, title)

    doc = Document()

    heading = title or f"{project_name} - Script"
    doc.add_heading(heading, level=0)
    doc.add_paragraph(f"Project: {project_name}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    for paragraph in script.split("\n"):
        if paragraph.strip():
            p = doc.add_paragraph(paragraph)
            p.style.font.size = Pt(11)

    safe_name = "".join(c if c.isalnum() or c in " -_" else "" for c in project_name)
    filename = f"{safe_name}_Script.docx"
    filepath = EXPORTS_DIR / filename

    doc.save(str(filepath))
    return str(filepath)


def export_full_report_txt(project_name: str, project_data: Dict,
                           step_outputs: Dict[int, str]) -> str:
    """Fallback: Export report as plain text if python-docx unavailable."""
    safe_name = "".join(c if c.isalnum() or c in " -_" else "" for c in project_name)
    filename = f"{safe_name}_Full_Report.txt"
    filepath = EXPORTS_DIR / filename

    with open(filepath, "w", encoding="utf-8") as f:
        f.write("YOUTUBE SEMANTIC OPTIMIZER - FULL REPORT\n")
        f.write("=" * 60 + "\n\n")
        f.write(f"Project: {project_name}\n")
        f.write(f"Topic: {project_data.get('topic', '')}\n")
        f.write(f"Channel: {project_data.get('channel', '')}\n")
        f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write("=" * 60 + "\n\n")

        for step_num in sorted(step_outputs.keys()):
            output = step_outputs[step_num]
            if not output:
                continue
            step_name = STEPS.get(step_num, f"Step {step_num}")
            f.write(f"\n{'=' * 60}\n")
            f.write(f"STEP {step_num}: {step_name}\n")
            f.write(f"{'=' * 60}\n\n")
            f.write(output)
            f.write("\n\n")

    return str(filepath)


def export_script_txt(project_name: str, script: str,
                      title: str = "") -> str:
    """Fallback: Export script as plain text."""
    safe_name = "".join(c if c.isalnum() or c in " -_" else "" for c in project_name)
    filename = f"{safe_name}_Script.txt"
    filepath = EXPORTS_DIR / filename

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f"SCRIPT: {title or project_name}\n")
        f.write("=" * 60 + "\n\n")
        f.write(script)

    return str(filepath)
